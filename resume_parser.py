"""Extract raw text from uploaded PDF and DOCX resumes (in memory, nothing is saved)."""
from pypdf import PdfReader
from docx import Document


def extract_pdf_text(uploaded_file):
    reader = PdfReader(uploaded_file)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_docx_text(uploaded_file):
    doc = Document(uploaded_file)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes keep skills inside tables, so read table cells too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_resume_text(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)
    if name.endswith(".docx"):
        return extract_docx_text(uploaded_file)
    raise ValueError("Only PDF and DOCX files are supported.")
